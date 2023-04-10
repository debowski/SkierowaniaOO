import tkinter as tk
# from tkinter import *
# from tkinter import StringVar
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import filedialog
from docxtpl import DocxTemplate
import docx
import pandas as pd
import os
import subprocess

szablon = "Szablony\\szablon.docx"
szablonWykaz = "Szablony\\szablonWykaz.docx"
tmp = "Szablony\\output1.docx"
domyslnyplik = "..\\Data\\WydrukiListXls.xlsx"

# ========= FUNKCJE ==========


def kofiguracja_domyslna(lab, lista_specjalnosci):
    if os.path.exists("..\\Data"):
        lab.configure(text=domyslnyplik)
        setZawody(domyslnyplik, lista_specjalnosci)
        return domyslnyplik


def wybPlik(lab, lista_specjalnosci) -> str:

    filetypes = (
        ('text files', '*.xlsx'),
        ('All files', '*.*')
    )

    plik = filedialog.askopenfilename(title='Wybierz plik',
                                      initialdir='..\\Data',
                                      filetypes=filetypes)
    lab.configure(text=plik)

    setZawody(plik, lista_specjalnosci)
    return plik


def otwarcie_folderu_wykazy():
    path = r"..\Data\Wykazy"
    subprocess.Popen(f'explorer "{path}"')


def otwarcie_folderu_skierowania():
    path = r"..\Data\Skierowania"
    subprocess.Popen(f'explorer "{path}"')


def setZawody(plik, lista_specjalnosci):
    df = pd.read_excel(plik)
    wynik = set(df['Specjalność/Zawód'].tolist())
    values = ['']
    for w in wynik:
        values += (w,)
    lista_specjalnosci['values'] = sorted(values)


def podajPlik(lab) -> str:
    return lab.cget("text")


def pobranieDatyWystawienia(dw) -> str:
    return dw.entry.get()


def pobranieDatyRozpoczecia(dr) -> str:
    return dr.entry.get()


def pobranieDatyZakonczenia(dz) -> str:
    return dz.entry.get()


def pobranieGodziny(godzina_spinbox, minuta_spinbox) -> str:
    godzina = godzina_spinbox.get()
    minuta = minuta_spinbox.get()
    return godzina+":"+minuta


def daneOddzialu(pole_dane_oddzialu) -> str:
    return pole_dane_oddzialu.get()


def ustawSpecjalnosc(spec) -> str:
    return spec.get()


def symbolZawodu(specjalnosc) -> str:

    zawody_dict = {
        'Cukiernik': '751201',
        'Fryzjer': '514101',
        'Sprzedawca': '522301',
        'Mechanik pojazdów samochodowych': '723103',
        'Kucharz': '512001',
        'Blacharz samochodowy': '721306',
        'Piekarz': '751204',
        'Stolarz': '752205',
        'Lakiernik': '713201',
        'Monter sieci, instalacji i urządzeń sanitarnych': '712616',
        'Elektromechanik pojazdów samochodowych': '741203',
        'Elektryk': '741103',
        'Monter zabudowy i robót wykończeniowych w budownictwie': '712905',
        'Mechanik-monter maszyn i urządzeń': '723310',
        'Magazynier-logistyk': '432106',
        '': 'N/A'
    }
    return zawody_dict[specjalnosc]

# =================================================


def wygeneruj_dokument(lista, pole_dane_oddzialu, lista_specjalnosci, spec, szablon, dw, dr, dz, godzina_spinbox, minuta_spinbox, wynik) -> None:
    # Otwórz plik xlsx
    df = pd.read_excel(open(lista, "rb"), dtype={'PESEL': str})

    filtered_df = df[df["Dane oddziału"].str.contains(daneOddzialu(pole_dane_oddzialu
                                                                   ), case=False) & df['Specjalność/Zawód'].str.contains(lista_specjalnosci.get(), case=False)]

    # .shape zwraca tupla wiersze, kolumny
    for linia in range(filtered_df.shape[0]):
        rekord = filtered_df.iloc[linia].to_dict()

        doc = DocxTemplate(szablon)

        context = {'dataWyst': pobranieDatyWystawienia(dw),
                   'imię': rekord['Imię'],
                   'nazwisko': rekord['Nazwisko'],
                   'dataUrodzenia': rekord['Data urodzenia'],
                   'miejsceUrodzenia': rekord['Miejsce urodzenia'],
                   'PESEL': rekord['PESEL'],
                   'zawod': lista_specjalnosci.get(),
                   'kodZawodu': symbolZawodu(lista_specjalnosci.get()),
                   'dataRozp': pobranieDatyRozpoczecia(dr),
                   'dataZako': pobranieDatyZakonczenia(dz),
                   'godzRozp': pobranieGodziny(godzina_spinbox, minuta_spinbox),
                   'PESEL': rekord['PESEL'],
                   'stopien': daneOddzialu(pole_dane_oddzialu)
                   }

        # print(context)

        # renderowane dokumentu (podstawianie danych ze słownika)
        doc.render(context)

        if not os.path.exists("..\\Data\\Skierowania"):
            os.mkdir("..\\Data\\Skierowania")

        # zapisywanie dokumentu
        doc.save("..\\Data\\Skierowania\\"+rekord['Dane oddziału']+rekord['Specjalność/Zawód'] +
                 rekord['Imię']+rekord['Nazwisko'] + ".docx")

        # informacja zwrotna
        wynik.configure(text=f"utworzono: {str(linia + 1)} dokumentów")


# *********************************************


# ==============================================

# generowanie wykazu uczniów skierowanych


def tworzenieWykazu(szablonWykaz, lista, tmp, wynik, pole_dane_oddzialu, lista_specjalnosci, spec, dw, dr, dz, godzina_spinbox, minuta_spinbox) -> None:

    # wstawianie listy uczniów na końcu dokumentu
    docTempl = docx.Document(szablonWykaz)
    dfw = pd.read_excel(open(lista, "rb"), dtype={'PESEL': str})
    filtered_dfw = dfw[dfw["Dane oddziału"].str.contains(daneOddzialu(pole_dane_oddzialu
                                                                      ), case=False) & dfw['Specjalność/Zawód'].str.contains(lista_specjalnosci.get(), case=False)]

    for linia in range(filtered_dfw.shape[0]):
        rekord = filtered_dfw.iloc[linia].to_dict()
        # new_paragraph = docTempl.add_paragraph(str(linia+1) + ". " + rekord['Imię']+ " " + rekord['Nazwisko'] )
        num_of_paragraphs = len(docTempl.paragraphs)
        # print(num_of_paragraphs)

        npar = docTempl.paragraphs[num_of_paragraphs-3]
        # print(npar)
        npar.add_run(
            (
                (
                    ((f"{str(linia + 1)}. " + rekord['Imię']) + " ")
                    + rekord['Nazwisko']
                )
                + "\n"
            )
        )

    if os.path.exists(tmp):
        os.remove(tmp)
    docTempl.save(tmp)

    # wstawienie danych do szablonu wykazu uczniów

    szablon = DocxTemplate(tmp)
    rekord = filtered_dfw.iloc[linia].to_dict()
    context = {'dataWyst': pobranieDatyWystawienia(dw),
               'imię': rekord['Imię'],
               'nazwisko': rekord['Nazwisko'],
               'dataUrodzenia': rekord['Data urodzenia'],
               'miejsceUrodzenia': rekord['Miejsce urodzenia'],
               'PESEL': rekord['PESEL'],
               'zawod': lista_specjalnosci.get(),
               'kodZawodu': symbolZawodu(lista_specjalnosci.get()),
               'dataRozp': pobranieDatyRozpoczecia(dr),
               'dataZako': pobranieDatyZakonczenia(dz),
               'godzRozp': pobranieGodziny(godzina_spinbox, minuta_spinbox),
               'PESEL': rekord['PESEL'],
               'stopien': daneOddzialu(pole_dane_oddzialu)
               }

    # renderowane dokumentu (podstawianie danych ze słownika)
    szablon.render(context)

    if not os.path.exists("..\\Data\\Wykazy"):
        os.mkdir("..\\Data\\Wykazy")

    # zapisywanie dokumentu
    szablon.save("..\\Data\\Wykazy\\"+rekord['Dane oddziału'] +
                 rekord['Specjalność/Zawód'] + ".docx")

    # informacja zwrotna

    wynik.configure(text=f"utworzono: {str(linia + 1)} pozycji")

# **********************************************


def wypiszDane(event, lista_specjalnosci, labprawa, plik, pole_dane_oddzialu):
    '''Wybranie osób z listy'''
    df = pd.read_excel(open(plik, "rb"))
    filtered_df = df[df["Dane oddziału"].str.contains(daneOddzialu(pole_dane_oddzialu
                                                                   ), case=False) & df['Specjalność/Zawód'].str.contains(lista_specjalnosci.get(), case=False)]
    tekst = ""
    numer = 1

    for linia in range(filtered_df.shape[0]):
        rekord = filtered_df.iloc[linia].to_dict()

        tekst = tekst + str(numer) + ". " + \
            rekord['Imię'] + " " + rekord['Nazwisko'] + "\n"

        numer = numer+1

    # wstawianie listy uczniów do ramki prawej
    labprawa.delete('1.0', END)
    labprawa.insert(tk.END, lista_specjalnosci.get() +
                    " ( " + pole_dane_oddzialu.get() + " ):\n---------------------\n")
    labprawa.insert(tk.END, tekst)

# ============================================================================


def mainapp():  # sourcery skip: extract-duplicate-method

    plik = ""

    root = ttk.Window(title="Skierowania 0.21",
                      themename="darkly", iconphoto="icon.png")
    root.resizable(False, False)
    # utworzenie ramki górnej
    frame_top = ttk.Frame(root)
    frame_top.grid(row=0, column=0, columnspan=2,
                   sticky='nsew', padx=10, pady=10)

    # utworzenie dwóch ramki w ramce górnej
    frame1 = ttk.Frame(frame_top)
    frame1.grid(row=0, column=0, sticky='nsew', padx=5, pady=5)

    frame2 = ttk.Frame(frame_top)
    frame2.grid(row=0, column=1, sticky='nsew', padx=10, pady=10)

    # ustawienie proporcji kolumn i wierszy
    root.columnconfigure(0, weight=1, minsize=200)
    root.columnconfigure(1, weight=1, minsize=200)

    frame1.columnconfigure(0, weight=1, minsize=100)
    frame1.columnconfigure(1, weight=1, minsize=100)
    frame2.columnconfigure(0, weight=1, minsize=100)
    frame2.columnconfigure(1, weight=1, minsize=100)

    # wiersz 0 - wybór pliku

    btn = ttk.Button(frame1, text="Wybierz plik",
                     command=lambda: wybPlik(lab, lista_specjalnosci))
    btn.grid(row=0, column=0, sticky='nsew', columnspan=2, padx=5, pady=5)

    # wiersz 1 - pokazanie wyboru pliku

    lab = ttk.Label(frame1, text="c:\\...",
                    bootstyle="inverse-success", foreground="black")
    lab.grid(row=1, column=0, columnspan=2, sticky='nsew', padx=5, pady=5)

    # wiersz 2 - poziom turnusu / klasy
    etykieta_dane_oddzialu = ttk.Label(frame1, text="Dane oddziału (klasa):")
    etykieta_dane_oddzialu.grid(
        row=2, column=0, sticky="snwe", padx=5,  pady=5)

    # pole_dane_oddzialu = ttk.Entry(frame1, justify="center")
    pole_dane_oddzialu = ttk.Spinbox(frame1, from_=1, to=3, justify="center")
    pole_dane_oddzialu.delete(0, END)
    pole_dane_oddzialu.insert(0, "1")

    pole_dane_oddzialu.bind("<ButtonRelease-1>", lambda event: wypiszDane(
        pole_dane_oddzialu.get(), lista_specjalnosci, labprawa, podajPlik(lab), pole_dane_oddzialu))

    pole_dane_oddzialu.grid(row=2, column=1, sticky="snwe", padx=5,  pady=5)

    # wiersz 3- specjalność
    etykieta_specjalnosc = ttk.Label(frame1, text="Specjalność:")
    etykieta_specjalnosc.grid(row=3, column=0, sticky="snwe", padx=5,  pady=5)

    spec = tk.StringVar(value="cukiernik")  # nie wiadomo czy to jest potrzebne

    current_var = tk.StringVar()
    lista_specjalnosci = ttk.Combobox(
        frame1, textvariable=current_var)
    lista_specjalnosci['values'] = ('Wybierz-plik-z-danymi')
    lista_specjalnosci['state'] = 'readonly'
    lista_specjalnosci.grid(row=3, column=1, sticky="snwe", padx=5,  pady=5)

    # lista_specjalnosci.bind("<<ComboboxSelected>>", lambda event : wypiszDane(lista_specjalnosci.get(), lista_specjalnosci, labprawa, podajPlik(lab), ))
    lista_specjalnosci.bind("<<ComboboxSelected>>", lambda event: wypiszDane(
        lista_specjalnosci.get(), lista_specjalnosci, labprawa, podajPlik(lab), pole_dane_oddzialu))

    # wiersz 4 - data wystawienia

    etykieta_data_wystawienia = ttk.Label(frame1, text="Data wystawienia:")
    etykieta_data_wystawienia.grid(
        row=4, column=0, sticky="snwe", padx=5,  pady=5)
    dw = ttk.DateEntry(frame1, firstweekday=0)
    dw.grid(row=4, column=1, sticky="snwe", padx=5,  pady=5)

    # wiersz 5 - data rozpoczęcia
    etykieta_data_reozpoczecia = ttk.Label(
        frame1, text="Data rozpoczęcia turnusu:")
    etykieta_data_reozpoczecia.grid(
        row=5, column=0, sticky="snwe", padx=5,  pady=5)
    dr = ttk.DateEntry(frame1, firstweekday=0)
    dr.grid(row=5, column=1, sticky="snwe", padx=5,  pady=5)

    # wiersz 6 - data zakończenia
    etykieta_data_zakonczenia = ttk.Label(
        frame1, text="Data zakończenia turnusu:")
    etykieta_data_zakonczenia.grid(
        row=6, column=0, sticky="snwe", padx=5,  pady=5)
    dz = ttk.DateEntry(frame1, firstweekday=0)
    dz.grid(row=6, column=1, sticky="snwe", padx=5,  pady=5)

    # wiersz 7 - godzina rozpoczęcia
    godzina_spinbox = ttk.Spinbox(
        frame1, from_=0, to=23, justify="center", format="%02.0f")
    minuta_spinbox = ttk.Spinbox(
        frame1, from_=0, to=59, justify="center", format="%02.0f")

    godzina_spinbox.delete(0, END)
    godzina_spinbox.insert(0, "8")
    minuta_spinbox.delete(0, END)
    minuta_spinbox.insert(0, "00")

    godzina_spinbox.grid(row=7, column=0,  sticky="snwe", padx=5, pady=5)
    minuta_spinbox.grid(row=7, column=1,  sticky="snwe", padx=5, pady=5)

    # wiersz 8 - przyciski do generowania
    przycisk_generuj_wykaz = ttk.Button(
        frame1, text="Wygeneruj wykaz", command=lambda: tworzenieWykazu(szablonWykaz, podajPlik(lab), tmp, wynik, pole_dane_oddzialu, lista_specjalnosci, spec, dw, dr, dz, godzina_spinbox, minuta_spinbox))
    przycisk_generuj_wykaz.grid(
        row=8, column=0, sticky="snwe", padx=5,  pady=5)

    przycisk_generuj = ttk.Button(
        frame1, text="Wygeneruj skierowania", command=lambda: wygeneruj_dokument(podajPlik(lab), pole_dane_oddzialu, lista_specjalnosci, spec, szablon, dw, dr, dz, godzina_spinbox, minuta_spinbox, wynik))
    przycisk_generuj.grid(row=8, column=1, sticky="snwe", padx=5,  pady=5)

    # wiersz 9 - ramka na status
    wynik = ttk.Label(frame1, text='Wynik',
                      justify="center", bootstyle="inverse-success", foreground="black")

    wynik.grid(row=9, column=0, columnspan=2, sticky='nsew', padx=5, pady=5)

    # zawartość ramki prawej
    labprawa = tk.Text(frame2, height=26, width=50)
    labprawa.grid(row=0, column=0, columnspan=2, sticky='nswe')

    scrollbar = ttk.Scrollbar(frame2, orient='vertical')
    scrollbar.grid(row=0, column=1, sticky='nse')

    scrollbar.config(command=labprawa.yview)
    labprawa.config(yscrollcommand=scrollbar.set)

    instrukcja = "[1] wybierz plik xlsx z danymi uczniów\n[2] wybierz klasę \n[3] Wybierz zawód z listy \n[4] Wskaż daty wystawienia, rozpoczęcia i zakończenia \n[5] Ustaw godzinę rozpoczęcia turnusu \n[6] wygeneruj wykaz oraz skierowania\n"

    labprawa.delete('1.0', END)
    labprawa.insert(tk.END, instrukcja)

    btnEkspWykazy = ttk.Button(
        frame1, text='Otwórz wykazy', command=otwarcie_folderu_wykazy)
    btnEkspWykazy.grid(row=10, column=0, sticky="snwe", padx=5, pady=5)

    btnEkspSkierowania = ttk.Button(
        frame1, text='Otwórz skierowania', command=otwarcie_folderu_skierowania)
    btnEkspSkierowania.grid(row=10, column=1, sticky="snwe", padx=5, pady=5)

    # załadowanie domyslnych ustawień
    kofiguracja_domyslna(lab, lista_specjalnosci)

    root.mainloop()


if __name__ == '__main__':
    mainapp()
